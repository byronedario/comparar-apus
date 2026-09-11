"""Lector de la hoja de presupuesto (tabla de rubros con cantidad y precio).

El presupuesto sirve para dos verificaciones que el APU por si solo no permite:

1. Que la cantidad de obra de cada rubro (los m2, m3, u que se van a ejecutar)
   sea la del presupuesto referencial. Bajarla cambia el monto del contrato sin
   tocar ningun APU, y ningun analisis de APUs se entera.
2. Que el valor unitario que el oferente puso en su presupuesto sea el mismo que
   sale de su propio APU. Si no cuadran, uno de los dos documentos no es el que
   va a ejecutar.

La cabecera se localiza por el NOMBRE de cada columna, no por su posicion, porque
cada entidad arma la tabla a su manera: "DESCRIPCION DEL RUBRO" frente a
"RUBRO / D E S C R I P C I O N", "PRECIO UNITARIO DEL RUBRO" frente a "PRECIO
UNITARIO OFERTADO". Las filas de capitulo ("OBRAS PRELIMINARES", "BIBLIOTECA")
no llevan cantidad y se descartan solas.
"""
import re

from apu import clean, num, grid, sin_tildes

# como se llama cada columna en las plantillas vistas
COLUMNAS = {
    'n':      (r'^ITEM$', r'^ITEM ?N', r'^N[°º]?$', r'^NO\.?$', r'^NUMERO',
               r'^RUBRO ?N'),
    'codigo': (r'^CODIGO', r'^COD'),
    'desc':   (r'^DESCRIPCION', r'^RUBRO', r'^DETALLE'),
    'uni':    (r'^UNIDAD', r'^UND$', r'^U\.?$'),
    'cant':   (r'^CANTIDAD', r'^CANT'),
    # "PRECIO UNITARIO" y "PRECIO TOTAL" empiezan igual: el unitario exige la
    # palabra completa para no quedarse tambien con la columna del total
    'punit':  (r'^PREC\w* ?UNITARIO', r'^P\.? ?UNITARIO', r'^V\.? ?UNITARIO',
               r'^PRECIO$'),
    'total':  (r'^PRECIO ?GLOBAL', r'^SUBTOTAL', r'^TOTAL', r'^VALOR ?TOTAL',
               r'^PREC\w* ?TOTAL'),
}


def _cual(texto):
    t = re.sub(r'\s+', ' ', sin_tildes(clean(texto))).strip()
    if not t:
        return None
    # "DESCRIPCION DEL CPC" y la columna de codigo no son la descripcion del rubro
    if re.search(r'\bCPC\b', t):
        return None
    for campo, patrones in COLUMNAS.items():
        for p in patrones:
            # el patron NO pasa por sin_tildes: .upper() convertiria \w en \W
            # y la columna dejaria de reconocerse en silencio
            if re.match(p, t):
                return campo
    return None


def parse(ws, max_fil=900, max_col=20):
    """{numero_de_item: {codigo, desc, uni, cant, punit, total}}.

    Si la hoja no parece un presupuesto devuelve {} en vez de reventar, para que
    el reporte se genere igual sin las columnas de presupuesto.
    """
    g = grid(ws, max_fil, max_col)

    # fila de encabezados: la que mapea mas columnas conocidas
    cab, fila_cab = {}, None
    for r in range(1, min(60, max_fil) + 1):
        m = {}
        for c in range(1, max_col + 1):
            campo = _cual(g[r][c])
            if campo and campo not in m:
                m[campo] = c
        if 'cant' in m and ('desc' in m or 'n' in m) and len(m) > len(cab):
            cab, fila_cab = m, r
    if not cab or 'cant' not in cab:
        return {}

    items, auto = {}, 0
    for r in range(fila_cab + 1, max_fil + 1):
        cant = num(g[r][cab['cant']])
        desc = clean(g[r][cab['desc']]) if 'desc' in cab else ''
        if cant is None:
            continue
        if 'desc' in cab and not desc:
            continue
        n = None
        if 'n' in cab:
            t = clean(g[r][cab['n']])
            if re.fullmatch(r'\d+', t):
                n = int(t)
        if n is None:                      # sin columna de item: se numera en orden
            auto += 1
            n = auto
        items[n] = {
            'codigo': clean(g[r][cab['codigo']]) if 'codigo' in cab else '',
            'desc': desc,
            'uni': clean(g[r][cab['uni']]) if 'uni' in cab else '',
            'cant': cant,
            'punit': num(g[r][cab['punit']]) if 'punit' in cab else None,
            'total': num(g[r][cab['total']]) if 'total' in cab else None,
            'hoja': ws.title, 'fila': r,
        }
    return items


def buscar_hoja(ruta):
    """Nombre de la hoja que parece el presupuesto, o None.

    Se prueba primero por nombre y, si eso no basta, por contenido: gana la hoja
    que produzca mas items.
    """
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    candidatas = [h for h in wb.sheetnames
                  if re.search(r'PRESUP|PRES[_\s-]|OFERTA|TABLA ?DE ?CANT',
                               sin_tildes(h))]
    if not candidatas:
        candidatas = [h for h in wb.sheetnames
                      if not clean(h).isdigit()
                      and not clean(h).upper().startswith('RUBRO')]
    mejor, punt = None, 0
    for h in candidatas:
        try:
            n = len(parse(wb[h]))
        except Exception:
            continue
        if n > punt:
            mejor, punt = h, n
    return mejor


def leer(ruta, hoja=None):
    """({n: item}, nombre_de_hoja). Autodetecta la hoja si no se indica.

    Si la ruta es un PDF, la tabla se saca con `tabla_pdf`, que lee por celdas
    y comprueba cada fila con su propia aritmetica. El presupuesto en PDF es la
    fuente mas fragil que maneja la skill: si existe el mismo dato en Excel, usa
    el Excel.
    """
    import openpyxl
    if re.search(r'\.pdf$', str(ruta), re.I):
        import tabla_pdf
        return tabla_pdf.leer(ruta), 'PDF'
    if not re.search(r'\.xls[xm]?$', str(ruta), re.I):
        return {}, None
    hoja = hoja or buscar_hoja(ruta)
    if not hoja:
        return {}, None
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    return parse(wb[hoja]), hoja
