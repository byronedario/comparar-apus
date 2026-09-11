"""Escribe el Excel del reporte, con las hojas enlazadas por formulas.

Por que formulas y no valores fijos: el revisor y su equipo repasan la hoja
DIFERENCIAS y reclasifican a mano lo que para ellos no es un error de fondo
(por ejemplo CEMENTO frente a CEMENTO PORTLAND TIPO 1). Si los conteos del
RESUMEN estuvieran escritos como numeros, esa decision no se reflejaria en
ninguna parte y habria que regenerar el archivo. Con formulas, al cambiar el
TIPO REVISADO se recalculan los conteos, el estado del rubro, los colores y la
hoja SOLO ERRORES.

La columna TIPO DETECTADO queda fija: es lo que encontro el analisis y sirve de
respaldo para saber que se cambio a mano. Lo que el equipo edita es TIPO
REVISADO.

Funciona en cualquier version de Excel: las hojas se enlazan con COUNTIFS e
INDEX/MATCH, no con formulas de matriz dinamica. `apu.reporte_estatico()` queda
para cuando se quiera un archivo sin ninguna formula.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TIPOS = ['DIFERENCIA', 'MENOR', 'ORDEN', 'TEXTO CORTADO']

TH = Font(bold=True, color='FFFFFF', size=10)
FH = PatternFill('solid', fgColor='1F3864')
ROJO = PatternFill('solid', fgColor='FFFFC7CE')
AMAR = PatternFill('solid', fgColor='FFFFF2CC')
VERD = PatternFill('solid', fgColor='FFC6EFCE')
AZUL = PatternFill('solid', fgColor='FFDDEBF7')
GRIS = PatternFill('solid', fgColor='FFF2F2F2')


def _cf(color):
    """Relleno para formato condicional.

    En un formato diferencial Excel pinta el fondo con bgColor; el fgColor que
    sirve en una celda normal ahi se ignora y la regla se aplica sin que se vea
    nada. Ademas el color va con alfa explicito: openpyxl antepone 00 a un color
    de seis digitos, y eso es transparente.
    """
    return PatternFill(patternType='solid', bgColor='FF' + color)


CF_ROJO, CF_AMAR, CF_VERD, CF_AZUL = (_cf('FFC7CE'), _cf('FFF2CC'),
                                      _cf('C6EFCE'), _cf('DDEBF7'))
BORDE = Border(*[Side(style='thin', color='BFBFBF')] * 4)

# que color le toca a cada valor. "MAYOR" es amarillo y no rojo por el criterio
# de fondo: que el oferente ponga mas cantidad no perjudica al contrato.
REGLAS = [
    (['DIFERENCIA', 'REVISAR'], CF_ROJO),
    (['MENOR', 'OBSERVACION MENOR', 'MAYOR'], CF_AMAR),
    (['ORDEN', 'TEXTO CORTADO', 'REDONDEO', 'ORDEN / TEXTO CORTADO'], CF_AZUL),
    (['OK'], CF_VERD),
]


def _hoja(ws, cabeceras, datos, anchos):
    ws.append(cabeceras)
    for c in range(1, len(cabeceras) + 1):
        cel = ws.cell(1, c)
        cel.font, cel.fill = TH, FH
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cel.border = BORDE
    for d in datos:
        ws.append(d)
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 32
    for fila in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2),
                             max_col=len(cabeceras)):
        for c in fila:
            c.border = BORDE
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.font = Font(size=10)


def _semaforo(ws, col, fil_ini, fil_fin):
    """Colorea una columna por su valor, con formato condicional.

    Tiene que ser condicional y no un relleno fijo: si el equipo reclasifica una
    fila, el color debe seguir a la nueva clasificacion.
    """
    L = get_column_letter(col)
    rango = '%s%d:%s%d' % (L, fil_ini, L, fil_fin)
    for valores, relleno in REGLAS:
        cond = ','.join('($%s%d="%s")' % (L, fil_ini, v) for v in valores)
        ws.conditional_formatting.add(
            rango, FormulaRule(formula=['OR(%s)' % cond], fill=relleno,
                               stopIfTrue=False))


def generar(salida, filas, resumen, filas_pres=(), presup_ref=None,
            presup_ofe=None, titulo_ref='', titulo_ofe='', correspondencia='',
            notas_extra=(), tolerancia_indirectos=5.0):
    """Escribe el Excel.

    filas        detalle de apu.comparar()
    resumen      una fila por rubro; el campo 12 es el precio unitario que sale
                 del APU de la oferta, para contrastarlo con su presupuesto
    filas_pres   hallazgos de apu.comparar_presupuesto(), o vacio
    """
    wb = Workbook()
    nfil = len(filas)
    ult = nfil + 1                       # ultima fila con datos en DIFERENCIAS

    # ---------------- DIFERENCIAS (la unica hoja que se edita) --------------
    ws = wb.create_sheet('DIFERENCIAS')
    _hoja(ws, ['RUBRO N.', 'CODIGO', 'RUBRO', 'SECCION', 'ITEM', 'CAMPO',
               'VALOR EN LA REFERENCIA', 'VALOR EN LA OFERTA',
               'TIPO DETECTADO', 'TIPO REVISADO', 'OBSERVACION'],
          [list(f) + [f[8], ''] for f in filas],
          [10, 10, 38, 14, 38, 26, 40, 40, 15, 15, 34])
    for r in range(2, ult + 1):
        ws.cell(r, 9).fill = GRIS        # detectado: no se toca
    if nfil:
        # showDropDown va invertido en el formato de Excel: False = mostrar la
        # flechita en la celda, que es lo que queremos
        dv = DataValidation(type='list', formula1='"%s"' % ','.join(TIPOS),
                            allow_blank=True, showDropDown=False)
        dv.error = 'Usa uno de los tipos de la lista.'
        dv.promptTitle = 'Reclasificar'
        dv.prompt = ('Cambia el tipo si el equipo decide que la diferencia no es '
                     'de fondo. El RESUMEN y SOLO ERRORES se actualizan solos.')
        ws.add_data_validation(dv)
        dv.add('J2:J%d' % ult)
        _semaforo(ws, 9, 2, ult)
        _semaforo(ws, 10, 2, ult)
        # columna auxiliar oculta: numera los graves para que SOLO ERRORES los
        # recupere con INDEX/MATCH en vez de con FILTER, que es formula de
        # matriz dinamica y openpyxl no la sabe escribir -- Excel la borra al
        # abrir el archivo y avisa de que "reparo" el libro
        ws.cell(1, 12).value = 'N GRAVE'
        for r in range(2, ult + 1):
            ws.cell(r, 12).value = ('=IF($J{0}="DIFERENCIA",'
                                    'COUNTIF($J$2:$J{0},"DIFERENCIA"),"")').format(r)
        ws.column_dimensions['L'].hidden = True
    ws.auto_filter.ref = 'A1:K%d' % max(ult, 1)

    # ---------------- RESUMEN ----------------------------------------------
    rs = wb.create_sheet('RESUMEN', 0)
    datos = []
    for r in resumen:
        n = r[0]
        pr = (presup_ref or {}).get(n, {})
        po = (presup_ofe or {}).get(n, {})
        datos.append(list(r[:6]) + [pr.get('cant'), po.get('cant'), None,
                                    r[11] if len(r) > 11 else None,
                                    po.get('punit'), None,
                                    None, None, None, None])
    _hoja(rs, ['RUBRO N.', 'CODIGO', 'RUBRO EN LA REFERENCIA',
               'RUBRO EN LA OFERTA', 'UNIDAD REFERENCIA', 'UNIDAD OFERTA',
               'CANT. PRESUP. REFERENCIA', 'CANT. PRESUP. OFERTA', 'CANTIDAD',
               'VALOR SEGUN APU', 'VALOR SEGUN PRESUPUESTO', 'VALOR OFERTADO',
               'GRAVES', 'ORDEN / CORTE', 'MENORES', 'ESTADO'],
          datos, [10, 10, 40, 40, 13, 12, 15, 14, 12, 13, 15, 12, 10, 12, 11, 22])

    for i in range(len(datos)):
        f = i + 2
        # cantidad de obra: de menos es observable, de mas no. Mismo criterio
        # que dentro del APU
        rs.cell(f, 9).value = ('=IF(OR($G{0}="",$H{0}=""),"",'
                               'IF(ABS($G{0}-$H{0})<0.0001,"OK",'
                               'IF($H{0}<$G{0},"REVISAR","MAYOR")))').format(f)
        # el APU del oferente y su propio presupuesto tienen que dar lo mismo
        rs.cell(f, 12).value = ('=IF(OR($J{0}="",$K{0}=""),"",'
                                'IF(ABS($J{0}-$K{0})<0.005,"OK","REVISAR"))').format(f)
        base = 'DIFERENCIAS!$A$2:$A$%d,$A%d,DIFERENCIAS!$J$2:$J$%d,' % (ult, f, ult)
        rs.cell(f, 13).value = '=COUNTIFS(%s"DIFERENCIA")' % base
        rs.cell(f, 14).value = ('=COUNTIFS({0}"ORDEN")+COUNTIFS({0}"TEXTO CORTADO")'
                                ).format(base)
        rs.cell(f, 15).value = '=COUNTIFS(%s"MENOR")' % base
        rs.cell(f, 16).value = (
            '=IF($M{0}>0,"REVISAR",IF($N{0}>0,"ORDEN / TEXTO CORTADO",'
            'IF($O{0}>0,"OBSERVACION MENOR","OK")))').format(f)
        for c in (7, 8, 10, 11):
            rs.cell(f, c).number_format = '#,##0.0000'
    if datos:
        fin = len(datos) + 1
        for col in (9, 12, 16):
            _semaforo(rs, col, 2, fin)
    rs.auto_filter.ref = 'A1:P%d' % max(len(datos) + 1, 1)

    # ---------------- PRESUPUESTO -------------------------------------------
    if filas_pres:
        pp = wb.create_sheet('PRESUPUESTO')
        _hoja(pp, ['RUBRO N.', 'RUBRO', 'CAMPO', 'VALOR EN LA REFERENCIA',
                   'VALOR EN LA OFERTA', 'TIPO'], [list(x) for x in filas_pres],
              [10, 52, 44, 38, 38, 16])
        _semaforo(pp, 6, 2, len(filas_pres) + 1)
        pp.auto_filter.ref = 'A1:F%d' % (len(filas_pres) + 1)

    # ---------------- SOLO ERRORES (vista viva) -----------------------------
    se = wb.create_sheet('SOLO ERRORES')
    _hoja(se, ['RUBRO N.', 'CODIGO', 'RUBRO', 'SECCION', 'ITEM', 'CAMPO',
               'VALOR EN LA REFERENCIA', 'VALOR EN LA OFERTA', 'TIPO DETECTADO',
               'TIPO REVISADO', 'OBSERVACION'], [],
          [10, 10, 38, 14, 38, 26, 40, 40, 15, 15, 34])
    if nfil:
        # Una fila por cada grave posible: el equipo puede subir a DIFERENCIA
        # una fila que hoy es menor, asi que se reserva sitio para todas.
        # INDEX/MATCH sobre la columna auxiliar funciona en cualquier version de
        # Excel y no necesita formulas de matriz.
        for i in range(nfil):
            f = i + 2
            for c in range(1, 12):
                L = get_column_letter(c)
                # el INDEX va dos veces a proposito: apuntando a una celda vacia
                # Excel devuelve 0, no vacio, y la columna OBSERVACION -- que
                # esta en blanco hasta que el equipo escriba algo -- se llenaria
                # de ceros. El IF de dentro los convierte en vacio de verdad.
                idx = ('INDEX(DIFERENCIAS!{L}$2:{L}${0},'
                       'MATCH(ROW()-1,DIFERENCIAS!$L$2:$L${0},0))').format(ult, L=L)
                se.cell(f, c).value = '=IFERROR(IF({0}="","",{0}),"")'.format(idx)
                se.cell(f, c).font = Font(size=10)
                se.cell(f, c).border = BORDE
                se.cell(f, c).alignment = Alignment(vertical='top', wrap_text=True)
        _semaforo(se, 10, 2, nfil + 1)
        se.auto_filter.ref = 'A1:K%d' % (nfil + 1)

    # ---------------- NOTAS -------------------------------------------------
    cuenta = lambda v: ('=COUNTIF(RESUMEN!$P$2:$P$%d,"%s")'
                        % (len(datos) + 1, v)) if datos else 0
    notas = [
        ['CRITERIO DE COMPARACION', ''],
        ['Referencia', titulo_ref],
        ['Documento analizado', titulo_ofe],
        ['Correspondencia', correspondencia],
        ['Campos comparados',
         'Nombre del rubro, unidad, indirectos y, en cada seccion (EQUIPOS / MANO '
         'DE OBRA / MATERIALES / TRANSPORTE): descripcion, cantidad y unidad de '
         'cada item. Si se leyeron los dos presupuestos, tambien la cantidad de '
         'cada rubro y el precio unitario contra el APU de la propia oferta.'],
        ['Campos NO comparados',
         'Tarifas, jornales, precios unitarios de insumos, rendimientos y costos: '
         'son propios de cada oferente. El monto de la oferta tampoco se observa, '
         'porque ya puntua por si mismo en la calificacion economica.'],
        ['Emparejamiento de items',
         'Los items se emparejan por descripcion, no por posicion. Un item listado '
         'en otro orden NO se reporta como faltante. "ITEM FALTA" y "ITEM ADICIONAL" '
         'solo aparecen cuando el item no esta en ninguna posicion de la seccion.'],
        ['', ''],
        ['COMO SE GRADUA CADA HALLAZGO', ''],
        ['El criterio',
         'La oferta no tiene que ser identica: tiene que cubrir el rubro. Lo que se '
         'observa es que FALTE algo o que venga de MENOS. Que sobre, no.'],
        ['DIFERENCIA (grave)',
         'Item que falta, cantidad menor a la de la referencia (aunque la baja sea '
         'de decimales o venga de redondear), cantidad que queda en cero, unidad '
         'cambiada, insumo sustituido, cambio de categoria salarial, indirectos '
         'fuera de la tolerancia y, en el presupuesto, cantidad de rubro menor o '
         'precio que no cuadra con el APU de la propia oferta.'],
        ['MENOR',
         'Item adicional, cantidad mayor a la de la referencia, mismo contenido '
         'escrito distinto e indirectos dentro de la tolerancia. Se anota, pero no '
         'es lo que hay que corregir: el oferente destina mas recursos al rubro.'],
        ['ORDEN', 'Los items estan en distinto orden. No es error por si mismo: '
                  'las cantidades y unidades igual se comparan contra el item correcto.'],
        ['TEXTO CORTADO', 'El texto coincide pero aparece truncado en el PDF porque '
                          'la celda impresa no alcanzo. Verificar en el original.'],
        ['Indirectos',
         'El oferente puede bajar sus indirectos hasta %g puntos porcentuales '
         'respecto a los de la referencia, pero no subirlos. Dentro de esa '
         'tolerancia el hallazgo es MENOR; fuera, grave. Se reporta agrupado, '
         'porque es una sola decision repetida en todos los rubros.'
         % tolerancia_indirectos],
        ['', ''],
        ['COMO USAR ESTE ARCHIVO', ''],
        ['Hoja DIFERENCIAS',
         'Es la unica que se edita. TIPO DETECTADO (gris) es lo que encontro el '
         'analisis y queda como respaldo. TIPO REVISADO es lo que el equipo decide, '
         'con lista desplegable. La columna OBSERVACION es para dejar el criterio.'],
        ['Al reclasificar una fila',
         'El RESUMEN recuenta, el ESTADO del rubro cambia, los colores siguen a la '
         'nueva clasificacion y la hoja SOLO ERRORES se rearma sola. No hace falta '
         'volver a generar el reporte.'],
        ['Hoja SOLO ERRORES',
         'Es una vista calculada de las filas con TIPO REVISADO = DIFERENCIA. No se '
         'edita: se cambia el tipo en DIFERENCIAS y esta hoja se actualiza.'],
        ['Requisito', 'Ninguno: funciona en cualquier version de Excel. Si '
                      'prefieres un archivo sin formulas, existe '
                      'apu.reporte_estatico().'],
        ['', ''],
        ['VERIFICACIONES DEL RESUMEN', ''],
        ['CANTIDAD',
         'Compara la cantidad de obra del presupuesto referencial contra la del '
         'presupuesto del oferente. Bajarla cambia el monto del contrato sin tocar '
         'ningun APU, por eso se revisa aparte. De menos es REVISAR; de mas, MAYOR.'],
        ['VALOR OFERTADO',
         'Compara el precio unitario que sale del APU del oferente contra el que '
         'consta en su propio presupuesto. Si no cuadran, la oferta es incoherente '
         'consigo misma y hay que pedir aclaracion.'],
        ['', ''],
        ['RESULTADO (se recalcula al reclasificar)', ''],
        ['Rubros con hallazgos graves', cuenta('REVISAR')],
        ['Rubros solo con observaciones menores', cuenta('OBSERVACION MENOR')],
        ['Rubros con orden distinto o texto cortado', cuenta('ORDEN / TEXTO CORTADO')],
        ['Rubros correctos', cuenta('OK')],
        ['Total de rubros revisados', len(resumen)],
        ['Hallazgos en el presupuesto',
         ('%d graves de %d' % (sum(1 for r in filas_pres if r[-1] == 'DIFERENCIA'),
                               len(filas_pres))) if filas_pres
         else 'no se compararon los presupuestos'],
    ]
    # notas_extra son pares [concepto, detalle]: un texto suelto se recorreria
    # letra por letra y llenaria la hoja de columnas
    notas += [list(x) if isinstance(x, (list, tuple)) else ['', str(x)]
              for x in notas_extra]
    _hoja(wb.create_sheet('NOTAS'), ['CONCEPTO', 'DETALLE'], notas, [40, 112])

    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']                  # RESUMEN ya se creo en la posicion 0
    wb.save(salida)
    return salida
